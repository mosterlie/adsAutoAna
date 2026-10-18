# -*- coding: utf-8 -*-
"""亚马逊 SP 批量操作（Bulk Operations）Excel 生成引擎。

移植自 adsHe/core/bulk_generator.py（已在真实上传中校准过 v8/v9 两次事故），
保持与亚马逊官方模板「商品推广活动」工作表 **31 列逐字一致**，仅额外追加一列
``ASIN`` 便于人工复核（亚马逊按表头精确匹配，未知列会被忽略）。

每个父体生成一个自动投放活动，行结构：

    Campaign -> [Bidding Adjustment(top) -> Bidding Adjustment(product page)]
    -> Ad Group -> Product Ad(每个子体一行) -> Product Targeting(4 组自动匹配)

上传前由 ``_validate_rows()`` 本地自检，任一不合规直接抛错、不产出文件，
避免「生成完上传才被整表拒绝」。

⚠️ 两条硬约束（来自历史事故，不要改动）：
1. 自动投放四组的实体层级是 ``Product Targeting``（不是 ``Targeting``），
   表达式写在官方列 **``拓展商品投放编号``** 里，取亚马逊受控裸值
   ``close-match`` / ``loose-match`` / ``substitutes`` / ``complements``；
2. 表头文字会被亚马逊精确匹配，改一个字等于整列失效（会报「以下列缺少值」）。
"""
from __future__ import annotations

import datetime
import io
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET_NAME = "Sponsored Products Campaigns"

# ---------------------------------------------------------------------------
# 亚马逊 SP 批量操作模板列顺序（与官方模板「商品推广活动」31 列逐字一致）
# ---------------------------------------------------------------------------
OFFICIAL_COLUMNS: List[str] = [
    "产品",              # 1  Product
    "实体层级",          # 2  Entity
    "操作",              # 3  Operation
    "广告活动编号",      # 4  Campaign ID
    "广告组编号",        # 5  Ad Group ID
    "广告组合编号",      # 6  Portfolio ID
    "广告编号",          # 7  Ad ID
    "关键词编号",        # 8  Keyword ID
    "商品投放 ID",       # 9  Product Targeting ID
    "广告活动名称",      # 10 Campaign Name
    "广告组名称",        # 11 Ad Group Name
    "开始日期",          # 12 Start Date
    "结束日期",          # 13 End Date
    "投放类型",          # 14 Targeting Type
    "状态",              # 15 State
    "每日预算",          # 16 Daily Budget
    "SKU",               # 17 SKU
    "广告组默认竞价",    # 18 Ad Group Default Bid
    "竞价",              # 19 Bid
    "关键词文本",        # 20 Keyword Text
    "母语关键词",        # 21 Native Language Keyword
    "母语区域",          # 22 Native Language Locale
    "匹配类型",          # 23 Match Type
    "竞价方案",          # 24 Bidding Strategy
    "广告位",            # 25 Placement
    "百分比",            # 26 Percentage
    "拓展商品投放编号",  # 27 Product Targeting Expression ★ 关键列
    "受众编号",          # 28 Audience ID
    "购物者群体占比",    # 29 Shopper Cohort Percentage
    "购物者群体类型",    # 30 Shopper Cohort Type
    "站点",              # 31 Sites
]

# 非官方辅助列：仅供人工复核，亚马逊会忽略未知列
HELPER_COLUMNS: List[str] = ["ASIN"]
BULK_COLUMNS: List[str] = OFFICIAL_COLUMNS + HELPER_COLUMNS

TARGETING_EXPR_COLUMN = "拓展商品投放编号"

VALID_ENTITIES = frozenset({
    "Campaign", "Ad Group", "Product Ad", "Keyword", "Negative Keyword",
    "Product Targeting", "Negative Product Targeting",
    "Campaign Negative Keyword", "Campaign Negative Product Targeting",
    "Bidding Adjustment",
})

AUTO_TARGETING = ["close-match", "loose-match", "substitutes", "complements"]

TARGETING_TYPE_AUTO = "Auto"
TARGETING_TYPE_MANUAL = "Manual"

PLACEMENT_TOP = "Placement Top"
PLACEMENT_PRODUCT_PAGE = "Placement Product Page"
PLACEMENT_REST_OF_SEARCH = "Placement Rest Of Search"
ALLOWED_PLACEMENTS = frozenset({PLACEMENT_TOP, PLACEMENT_PRODUCT_PAGE, PLACEMENT_REST_OF_SEARCH})

PERCENTAGE_MIN = 0
PERCENTAGE_MAX = 900

BIDDING_STRATEGIES = (
    "Fixed bid",
    "Dynamic bids - down only",
    "Dynamic bids - up and down",
)


def _to_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _txt(value: Any) -> str:
    return "" if value is None else str(value).strip()


def generate_bulk_excel(products: List[Dict[str, Any]],
                        settings: Dict[str, Any]) -> io.BytesIO:
    """按亚马逊官方格式组装 Sponsored Products 批量操作表格。

    products 每项:
        {parent_sku, product_name, assigned_budget,
         children: [{msku, asin}]}          # 只有 children 非空的项才会产出
    """
    settings = settings or {}
    rows: List[Dict[str, Any]] = []

    today = datetime.date.today()
    today_str = today.strftime("%Y%m%d")
    start_date: Any = settings.get("start_date") or today_str
    try:
        start_date = int(str(start_date).replace("-", "").strip())
    except (TypeError, ValueError):
        start_date = today_str

    bidding_strategy = settings.get("bidding_strategy") or "Fixed bid"
    if bidding_strategy not in BIDDING_STRATEGIES:
        bidding_strategy = "Fixed bid"

    default_bid = _to_float(settings.get("default_bid", 25.0), 25.0)
    top_adj = _to_float(settings.get("placement_top_pct", 0.0))
    prod_adj = _to_float(settings.get("placement_product_pct", 0.0))

    targeting_bids = {
        "close-match": _to_float(settings.get("bid_close", 25.0), 25.0),
        "loose-match": _to_float(settings.get("bid_loose", 20.0), 20.0),
        "substitutes": _to_float(settings.get("bid_substitutes", 25.0), 25.0),
        "complements": _to_float(settings.get("bid_complements", 20.0), 20.0),
    }
    include_targeting = bool(settings.get("include_targeting", True))
    include_placement = bool(settings.get("include_placement", True))
    # 编号列: True=用父 SKU 作为表内引用键(实盘验证过的写法); False=留空交给亚马逊分配
    fill_ids = bool(settings.get("fill_ids", True))
    suffix = _txt(settings.get("campaign_suffix")) or "自动"
    portfolio_id = _txt(settings.get("portfolio_id"))

    for item in products or []:
        parent_sku = _txt(item.get("parent_sku"))
        children = [c for c in (item.get("children") or []) if _txt(c.get("msku"))]
        if not parent_sku or not children:
            continue

        prod_name = _txt(item.get("product_name")) or "商品"
        budget = _to_float(
            item.get("assigned_budget",
                     item.get("custom_budget", item.get("suggested_budget", 100.0))),
            100.0)
        budget = max(1.0, round(budget, 2))

        campaign_id = parent_sku if fill_ids else ""
        ad_group_id = f"{parent_sku}_ag" if fill_ids else ""
        campaign_name = f"{parent_sku}_{prod_name}_{suffix}_{today_str}"
        ad_group_name = campaign_name

        # 1) Campaign
        rows.append({
            "产品": "Sponsored Products",
            "实体层级": "Campaign",
            "操作": "Create",
            "广告活动编号": campaign_id,
            "广告组合编号": portfolio_id,
            "广告活动名称": campaign_name,
            "开始日期": start_date,
            "投放类型": TARGETING_TYPE_AUTO,
            "状态": "enabled",
            "每日预算": budget,
            "竞价方案": bidding_strategy,
        })

        # 2) Bidding Adjustment（百分比 0 = 空操作，跳过）
        if include_placement:
            for placement, pct in ((PLACEMENT_TOP, top_adj),
                                   (PLACEMENT_PRODUCT_PAGE, prod_adj)):
                if pct <= 0:
                    continue
                rows.append({
                    "产品": "Sponsored Products",
                    "实体层级": "Bidding Adjustment",
                    "操作": "Create",
                    "广告活动编号": campaign_id,
                    "广告活动名称": campaign_name,
                    "竞价方案": bidding_strategy,
                    "广告位": placement,
                    "百分比": int(pct) if float(pct).is_integer() else round(pct, 2),
                })

        # 3) Ad Group
        rows.append({
            "产品": "Sponsored Products",
            "实体层级": "Ad Group",
            "操作": "Create",
            "广告活动编号": campaign_id,
            "广告组编号": ad_group_id,
            "广告活动名称": campaign_name,
            "广告组名称": ad_group_name,
            "状态": "enabled",
            "广告组默认竞价": round(default_bid, 2),
        })

        # 4) Product Ad（每个子体一行）
        for child in children:
            rows.append({
                "产品": "Sponsored Products",
                "实体层级": "Product Ad",
                "操作": "Create",
                "广告活动编号": campaign_id,
                "广告组编号": ad_group_id,
                "广告活动名称": campaign_name,
                "广告组名称": ad_group_name,
                "SKU": _txt(child.get("msku")),
                "ASIN": _txt(child.get("asin")),
                "状态": "enabled",
            })

        # 5) Product Targeting（4 组自动匹配）
        if include_targeting:
            for pred in AUTO_TARGETING:
                rows.append({
                    "产品": "Sponsored Products",
                    "实体层级": "Product Targeting",
                    "操作": "Create",
                    "广告活动编号": campaign_id,
                    "广告组编号": ad_group_id,
                    "广告活动名称": campaign_name,
                    "广告组名称": ad_group_name,
                    "状态": "enabled",
                    "竞价": round(targeting_bids[pred], 2),
                    TARGETING_EXPR_COLUMN: pred,
                })

    if not rows:
        raise ValueError("没有可生成的活动：请至少选择一个父体（且该父体下有带 SKU 的子体）")

    _validate_rows(rows)
    return _write_workbook(rows)


def _write_workbook(rows: List[Dict[str, Any]]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(BULK_COLUMNS)

    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="F2F4F7")
    for cell in ws[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")

    for r in rows:
        ws.append([r.get(c) for c in BULK_COLUMNS])

    # 关键列设成文本格式，避免亚马逊把 SKU / 编号当数字截断（如 0012 → 12）
    text_cols = {c for c in ("广告活动编号", "广告组编号", "广告组合编号",
                             "广告活动名称", "广告组名称", "SKU", "ASIN")}
    for idx, name in enumerate(BULK_COLUMNS, start=1):
        if name in text_cols:
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                for cell in row:
                    cell.number_format = "@"

    widths = {"产品": 18, "实体层级": 20, "广告活动名称": 46, "广告组名称": 46,
              "SKU": 30, "开始日期": 12, TARGETING_EXPR_COLUMN: 22}
    for idx, name in enumerate(BULK_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 14)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _validate_rows(rows: List[Dict[str, Any]]) -> None:
    """上传前自检：把亚马逊会拒绝的字段值挡在本地。"""
    problems: List[str] = []

    for i, r in enumerate(rows, start=2):
        entity = r.get("实体层级")
        where = f"第{i}行[{entity}]"

        if entity not in VALID_ENTITIES:
            problems.append(f"{where} 非法实体层级")
            continue

        if entity == "Campaign":
            if _txt(r.get("投放类型")) not in (TARGETING_TYPE_AUTO, TARGETING_TYPE_MANUAL):
                problems.append(f"{where} 投放类型必须是 Auto / Manual")
            if _to_float(r.get("每日预算")) <= 0:
                problems.append(f"{where} 每日预算必须大于 0")
            if _txt(r.get("竞价方案")) not in BIDDING_STRATEGIES:
                problems.append(f"{where} 竞价方案不在受控取值内：{r.get('竞价方案')!r}")
            if not _txt(r.get("广告活动名称")):
                problems.append(f"{where} 广告活动名称为空")

        elif entity == "Ad Group":
            if _to_float(r.get("广告组默认竞价")) <= 0:
                problems.append(f"{where} 广告组默认竞价必须大于 0")
            if not _txt(r.get("广告组名称")):
                problems.append(f"{where} 广告组名称为空")

        elif entity == "Product Ad":
            if not _txt(r.get("SKU")):
                problems.append(f"{where} SKU 为空（亚马逊按 SKU 挂载商品）")

        elif entity == "Bidding Adjustment":
            if _txt(r.get("广告位")) not in ALLOWED_PLACEMENTS:
                problems.append(f"{where} 非法广告位：{r.get('广告位')!r}")
            pct = _to_float(r.get("百分比"), -1)
            if not (PERCENTAGE_MIN <= pct <= PERCENTAGE_MAX):
                problems.append(f"{where} 百分比越界：{r.get('百分比')!r}（应为 0-900）")

        elif entity == "Product Targeting":
            expr = _txt(r.get(TARGETING_EXPR_COLUMN))
            if not expr:
                problems.append(
                    f"{where} 「{TARGETING_EXPR_COLUMN}」为空（亚马逊会报「以下列缺少值」）")
            elif expr not in AUTO_TARGETING:
                problems.append(
                    f"{where} 「{TARGETING_EXPR_COLUMN}」必须是裸受控值 "
                    f"{AUTO_TARGETING}，当前 {expr!r}")
            if _txt(r.get("关键词文本")) or _txt(r.get("匹配类型")):
                problems.append(f"{where} 投放行的「关键词文本」「匹配类型」必须留空")

    if problems:
        head = problems[:12]
        more = f"\n  … 另有 {len(problems) - len(head)} 处" if len(problems) > len(head) else ""
        raise ValueError("生成的 Bulk 表格未通过亚马逊规范自检：\n  - "
                         + "\n  - ".join(head) + more)
